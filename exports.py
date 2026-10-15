import io
from datetime import datetime
import pandas as pd


def show_aggrid(df, height=400):
    try:
        from st_aggrid import AgGrid, GridOptionsBuilder
        gb = GridOptionsBuilder.from_dataframe(df)
        gb.configure_default_column(filter=True, sortable=True, resizable=True)
        AgGrid(df, gridOptions=gb.build(), height=height, theme="balham_dark")
    except:
        import streamlit as st
        st.dataframe(df, width='stretch')


def generar_plantilla():
    buf = io.BytesIO()
    cols = [
        "fecha", "vehiculo", "conductor", "tipo_combustible",
        "kilometraje", "litros", "custo", "factura",
        "estacion", "vale", "latitud", "longitud"
    ]
    pd.DataFrame(columns=cols).to_excel(buf, index=False)
    buf.seek(0)
    return buf.getvalue()


def generar_excel(df, titulo="Reporte", tenant_id=1):
    from company import load_config
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from openpyxl import load_workbook

    cfg = load_config(tenant_id)
    buf = io.BytesIO()
    writer = pd.ExcelWriter(buf, engine="openpyxl")

    export_cols = [c for c in df.columns if c not in ['id', 'capturado_por', 'fecha_captura', 'tenant_id']]
    df_export = df[export_cols].copy()

    df_export.to_excel(writer, sheet_name="Datos", index=False, startrow=4)
    writer.close()

    buf.seek(0)
    wb = load_workbook(buf)
    ws = wb.active

    green_bg = PatternFill(start_color='0B6623', end_color='0B6623', fill_type='solid')
    light_green_bg = PatternFill(start_color='E8F5E9', end_color='E8F5E9', fill_type='solid')
    white_font = Font(name='Calibri', bold=True, color='FFFFFF', size=11)
    title_font = Font(name='Calibri', bold=True, size=18, color='0B6623')
    subtitle_font = Font(name='Calibri', size=10, color='666666')
    stat_font = Font(name='Calibri', bold=True, size=11, color='0B6623')
    data_font = Font(name='Calibri', size=10)
    thin_border = Border(
        left=Side(style='thin', color='CCCCCC'),
        right=Side(style='thin', color='CCCCCC'),
        top=Side(style='thin', color='CCCCCC'),
        bottom=Side(style='thin', color='CCCCCC')
    )

    empresa = cfg.get('empresa', 'CyberFuel')
    rfc = cfg.get('rfc', '')
    email = cfg.get('email', '')
    telefono = cfg.get('telefono', '')

    ws.merge_cells('A1:H1')
    ws['A1'] = f'  {empresa}'
    ws['A1'].font = title_font

    ws.merge_cells('A2:H2')
    info_parts = [f'RFC: {rfc}' if rfc else '', f'Tel: {telefono}' if telefono else '', f'Email: {email}' if email else '']
    info_line = '  |  '.join([p for p in info_parts if p])
    ws['A2'] = info_line
    ws['A2'].font = subtitle_font

    ws.merge_cells('A3:H3')
    ws['A3'] = f'{titulo}  -  Generado: {datetime.now().strftime("%d/%m/%Y %H:%M")}'
    ws['A3'].font = subtitle_font

    ws.merge_cells('A4:H4')
    total_gasto = df_export['custo'].sum() if 'custo' in df_export.columns else 0
    total_litros = df_export['litros'].sum() if 'litros' in df_export.columns else 0
    n_cargas = len(df_export)
    n_vehiculos = df_export['vehiculo'].nunique() if 'vehiculo' in df_export.columns else 0
    ws['A4'] = f'Registros: {n_cargas}  |  Vehiculos: {n_vehiculos}  |  Gasto Total: ${total_gasto:,.2f}  |  Litros Totales: {total_litros:,.1f}'
    ws['A4'].font = stat_font

    col_names = {}
    for col_idx in range(1, ws.max_column + 1):
        cell = ws.cell(row=5, column=col_idx)
        cell.font = white_font
        cell.fill = green_bg
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = thin_border
        col_names[col_idx] = str(cell.value) if cell.value else ''

    currency_cols = {'custo', 'custo_por_km'}
    number_cols = {'litros', 'kilometraje', 'rendimento', 'latitud', 'longitud', 'rendimiento_min', 'km_inicial', 'km_proximo_servicio', 'intervalo_km_servicio', 'capacidade'}

    for row_idx in range(6, ws.max_row + 1):
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = data_font
            cell.border = thin_border
            col_name = col_names.get(col_idx, '')

            if col_name in currency_cols:
                cell.number_format = '$#,##0.00'
                cell.alignment = Alignment(horizontal='right')
            elif col_name in number_cols:
                cell.number_format = '#,##0.00'
                cell.alignment = Alignment(horizontal='right')
            else:
                cell.alignment = Alignment(horizontal='left')

            if (row_idx - 6) % 2 == 0:
                cell.fill = light_green_bg

    for col_idx in range(1, ws.max_column + 1):
        max_length = len(str(col_names.get(col_idx, '')))
        for row_idx in range(6, min(ws.max_row + 1, 106)):
            cell = ws.cell(row=row_idx, column=col_idx)
            try:
                if cell.value is not None:
                    val_len = len(str(cell.value))
                    max_length = max(max_length, val_len)
            except:
                pass
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = min(max(max_length + 3, 10), 35)

    ws.freeze_panes = 'A6'
    ws.auto_filter.ref = f'A5:{get_column_letter(ws.max_column)}{ws.max_row}'

    buf2 = io.BytesIO()
    wb.save(buf2)
    buf2.seek(0)
    return buf2.getvalue()


def generar_pdf(df, titulo="Reporte", tenant_id=1):
    from company import load_config
    cfg = load_config(tenant_id)

    export_cols = [c for c in df.columns if c not in ['id', 'capturado_por', 'fecha_captura', 'tenant_id']]
    df_export = df[export_cols].copy()

    try:
        return _pdf_fpdf2(df_export, titulo, cfg)
    except ImportError:
        pass
    try:
        return _pdf_reportlab(df_export, titulo, cfg)
    except ImportError:
        pass
    return _pdf_text_fallback(df_export, titulo, cfg)


def _pdf_fpdf2(df, titulo, cfg):
    from fpdf import FPDF

    empresa = cfg.get('empresa', 'CyberFuel')
    rfc = cfg.get('rfc', '')
    telefono = cfg.get('telefono', '')
    email = cfg.get('email', '')

    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()

    pdf.set_fill_color(11, 102, 35)
    pdf.rect(10, 10, 190, 25, 'F')

    pdf.set_font('Helvetica', 'B', 20)
    pdf.set_text_color(255, 255, 255)
    pdf.set_xy(15, 12)
    pdf.cell(0, 10, f'  {empresa}', ln=True)

    pdf.set_font('Helvetica', '', 9)
    pdf.set_xy(15, 22)
    info_parts = [f'RFC: {rfc}' if rfc else '', f'Tel: {telefono}' if telefono else '', f'Email: {email}' if email else '']
    info_line = '  |  '.join([p for p in info_parts if p])
    pdf.cell(0, 8, f'  {info_line}')

    pdf.set_text_color(0, 0, 0)
    pdf.set_y(40)
    pdf.set_font('Helvetica', 'B', 14)
    pdf.cell(0, 8, titulo, ln=True)

    pdf.set_font('Helvetica', '', 9)
    pdf.set_text_color(100, 100, 100)

    total_gasto = df['custo'].sum() if 'custo' in df.columns else 0
    total_litros = df['litros'].sum() if 'litros' in df.columns else 0
    n_cargas = len(df)
    n_vehiculos = df['vehiculo'].nunique() if 'vehiculo' in df.columns else 0

    pdf.cell(0, 5, f'Generado: {datetime.now().strftime("%d/%m/%Y %H:%M")}  |  Registros: {n_cargas}  |  Vehiculos: {n_vehiculos}  |  Gasto: ${total_gasto:,.2f}  |  Litros: {total_litros:,.1f}', ln=True)
    pdf.ln(5)

    cols = list(df.columns)
    if not cols:
        pdf.set_font('Helvetica', '', 10)
        pdf.cell(0, 10, 'Sin datos', ln=True)
        buf = io.BytesIO()
        pdf.output(buf)
        buf.seek(0)
        return buf.getvalue()

    priority_cols = ['fecha', 'vehiculo', 'conductor', 'tipo_combustible', 'kilometraje', 'litros', 'custo', 'factura', 'estacion', 'rendimento', 'alerta_anomalia']
    display_cols = [c for c in priority_cols if c in cols]
    if not display_cols:
        display_cols = cols[:8]

    df_display = df[display_cols].copy()
    col_widths = _calc_col_widths(pdf, display_cols, df_display)

    pdf.set_font('Helvetica', 'B', 7)
    pdf.set_fill_color(11, 102, 35)
    pdf.set_text_color(255, 255, 255)

    for i, col in enumerate(display_cols):
        w = col_widths[i]
        pdf.cell(w, 7, _truncate(str(col).upper(), 12), border=1, fill=True, align='C')
    pdf.ln()

    pdf.set_font('Helvetica', '', 7)
    pdf.set_text_color(0, 0, 0)

    row_limit = min(len(df_display), 500)
    for row_idx in range(row_limit):
        if row_idx % 2 == 0:
            pdf.set_fill_color(232, 245, 233)
        else:
            pdf.set_fill_color(255, 255, 255)

        for i, col in enumerate(display_cols):
            w = col_widths[i]
            val = df_display.iloc[row_idx][col]
            val_str = _format_cell(val, col)
            pdf.cell(w, 6, _truncate(val_str, 18), border=1, fill=True, align='R' if col in ('custo', 'litros', 'kilometraje', 'rendimento', 'custo_por_km') else 'L')
        pdf.ln()

    if len(df_display) > row_limit:
        pdf.set_font('Helvetica', 'I', 8)
        pdf.set_text_color(100, 100, 100)
        pdf.cell(0, 8, f'... y {len(df_display) - row_limit} registros mas', ln=True)

    pdf.set_y(-20)
    pdf.set_font('Helvetica', 'I', 7)
    pdf.set_text_color(150, 150, 150)
    pdf.cell(0, 5, f'{empresa} | {titulo} | Pagina {pdf.page_no()}', align='C')

    buf = io.BytesIO()
    pdf.output(buf)
    buf.seek(0)
    return buf.getvalue()


def _pdf_reportlab(df, titulo, cfg):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer

    empresa = cfg.get('empresa', 'CyberFuel')
    rfc = cfg.get('rfc', '')

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(letter))
    styles = getSampleStyleSheet()
    elements = []

    elements.append(Paragraph(f'<font size="18" color="#0B6623"><b>{empresa}</b></font>', styles['Normal']))
    elements.append(Paragraph(f'<font size="9" color="#666666">RFC: {rfc} | {titulo} | {datetime.now().strftime("%d/%m/%Y %H:%M")}</font>', styles['Normal']))
    elements.append(Spacer(1, 12))

    priority_cols = ['fecha', 'vehiculo', 'conductor', 'tipo_combustible', 'kilometraje', 'litros', 'custo', 'factura', 'estacion', 'rendimento', 'alerta_anomalia']
    display_cols = [c for c in priority_cols if c in df.columns]
    if not display_cols:
        display_cols = list(df.columns)[:8]

    df_display = df[display_cols].head(200)
    data = [[str(c).upper() for c in display_cols]]
    for _, row in df_display.iterrows():
        data.append([_truncate(_format_cell(row[c], c), 18) for c in display_cols])

    table = Table(data, repeatRows=1)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0B6623')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 7),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.HexColor('#E8F5E9'), colors.white]),
        ('ALIGN', (0, 0), (-1, 0), 'CENTER'),
    ]))

    elements.append(table)
    doc.build(elements)
    buf.seek(0)
    return buf.getvalue()


def _pdf_text_fallback(df, titulo, cfg):
    empresa = cfg.get('empresa', 'CyberFuel')
    buf = io.BytesIO()
    lines = []
    lines.append('%PDF-1.4')
    lines.append('1 0 obj<</Type/Catalog/Pages 2 0 R>>endobj')
    lines.append('2 0 obj<</Type/Pages/Kids[3 0 R]/Count 1>>endobj')
    text_content = f'{empresa} - {titulo}\n'
    text_content += f'Generado: {datetime.now().strftime("%d/%m/%Y %H:%M")}\n\n'
    for _, row in df.head(100).iterrows():
        text_content += ' | '.join([str(v) for v in row.values]) + '\n'
    escaped = text_content.replace('\\', '\\\\').replace('(', '\\(').replace(')', '\\)')
    lines.append('3 0 obj<</Type/Page/Parent 2 0 R/MediaBox[0 0 612 792]/Contents 4 0 R>>endobj')
    lines.append(f'4 0 obj<</Length {len(escaped)}>>stream\nBT /F1 8 Tf 50 750 Td({escaped})Tj ET\nendstream\nendobj')
    lines.append('5 0 obj<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>endobj')
    lines.append('xref\n0 6\n0000000000 65535 f \ntrailer<</Root 1 0 R/Size 6>>startxref\n0\n%%EOF')
    content = '\n'.join(lines).encode('latin-1', errors='replace')
    buf.write(content)
    buf.seek(0)
    return buf.getvalue()


def _truncate(text, max_len=18):
    text = str(text) if text is not None else ''
    return text[:max_len] + '..' if len(text) > max_len else text


def _format_cell(val, col_name):
    if pd.isna(val):
        return ''
    try:
        if col_name in ('custo', 'custo_por_km'):
            return f'${float(val):,.2f}'
        if col_name in ('litros', 'kilometraje', 'rendimento'):
            return f'{float(val):,.2f}'
    except (ValueError, TypeError):
        pass
    return str(val) if val is not None else ''


def _calc_col_widths(pdf, cols, df):
    available = 190
    n = len(cols)
    if n == 0:
        return []
    base_widths = {
        'fecha': 20, 'vehiculo': 18, 'conductor': 22,
        'tipo_combustible': 18, 'kilometraje': 18, 'litros': 14,
        'custo': 16, 'factura': 18, 'estacion': 22,
        'vale': 14, 'rendimento': 14, 'alerta_anomalia': 16,
        'alerta_mantenimiento': 16, 'latitud': 18, 'longitud': 18,
    }
    widths = []
    for col in cols:
        w = base_widths.get(col, 18)
        widths.append(w)
    total = sum(widths)
    if total > 0:
        widths = [w * available / total for w in widths]
    return widths